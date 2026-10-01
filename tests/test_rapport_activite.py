# -*- coding: utf-8 -*-
"""
Tests de logic.rapports.generer_rapport_activite : rapport suivant la
structure institutionnelle (taux de réponses par pays, analyse des
reportings par pays, texte adaptatif selon le taux de transmission
réellement calculé).
"""
import os
import pandas as pd
import pytest

from logic.rapports import (
    generer_rapport_activite, _compter_participants_recus_par_pays,
    participants_attendus_par_pays, LIBELLES_STATUT,
)


def _evt(pays, fichier_source, survenu=True, statut=None, mois="Avril"):
    return {
        "reference": "RPSYGPS01", "evenement": "Test", "survenance": "Oui" if survenu else "Non",
        "occurrence": None, "motif": None, "statut": statut, "duree_indispo_min": None,
        "commentaire": None, "domaine": "SYGMA", "site": "PS",
        "fichier_source": fichier_source,
        "detection_confiance": 1.0, "detection_avertissement": None, "hors_referentiel": False,
        "survenance_bin": 1 if survenu else 0, "occurrence_num": None,
        "pays": pays, "annee": "2026", "mois": mois,
    }


class TestParticipantsAttendus:

    def test_total_correspond_au_document_de_reference(self, dossier_config_isole):
        """Le total des 6 pays doit correspondre exactement aux 60
        participants illustratifs de la configuration par défaut."""
        assert sum(participants_attendus_par_pays().values()) == 60

    def test_six_pays_couverts(self, dossier_config_isole):
        assert len(participants_attendus_par_pays()) == 6

    def test_valeurs_individuelles_correspondent_au_reference(self, dossier_config_isole):
        assert participants_attendus_par_pays()["Cameroun"] == 10
        assert participants_attendus_par_pays()["Congo"] == 10
        assert participants_attendus_par_pays()["Tchad"] == 10
        assert participants_attendus_par_pays()["Guinée Équatoriale"] == 10
        assert participants_attendus_par_pays()["Gabon"] == 10
        assert participants_attendus_par_pays()["Centrafrique"] == 10


class TestComptageParticipantsRecus:

    def test_compte_les_dossiers_banque_distincts_par_pays(self, dossier_config_isole):
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/AUTREBANK/f.xls"),
            _evt("Cameroun", "Cameroun/MOIS AVRIL 2026/BANQUE_X/f.xls"),
        ])
        recus = _compter_participants_recus_par_pays(evt_df)
        assert recus["Tchad"] == 2
        assert recus["Cameroun"] == 1

    def test_meme_banque_plusieurs_fichiers_comptee_une_fois(self, dossier_config_isole):
        """Une meme banque avec plusieurs fichiers (SYGMA + SWIFT par
        exemple) ne doit compter qu'une seule fois comme participant."""
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/sygma.xls"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/swift.xls"),
        ])
        recus = _compter_participants_recus_par_pays(evt_df)
        assert recus["Tchad"] == 1

    def test_pays_absent_du_lot_nest_pas_dans_le_resultat(self, dossier_config_isole):
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        recus = _compter_participants_recus_par_pays(evt_df)
        assert "Congo" not in recus


class TestReglagesParticipants:
    """
    Attendus et Recus sont reglables par l'agent depuis Parametres, et
    doivent s'appliquer automatiquement a la generation du rapport
    d'activite, sans qu'aucun code ne soit a modifier.
    """

    def test_attendus_par_defaut_sans_configuration(self, dossier_config_isole):
        from logic.rapports import participants_attendus_par_pays
        attendus = participants_attendus_par_pays()
        assert attendus["Cameroun"] == 10
        assert sum(attendus.values()) == 60

    def test_attendus_personnalises_par_lagent_sont_pris_en_compte(self, dossier_config_isole):
        from logic.rapports import participants_attendus_par_pays
        from logic.configuration import charger_configuration, sauvegarder_configuration

        config = charger_configuration()
        config["participants_attendus"] = {"Cameroun": 30, "Congo": 20}
        sauvegarder_configuration(config)

        attendus = participants_attendus_par_pays()
        assert attendus["Cameroun"] == 30
        assert attendus["Congo"] == 20

    def test_recu_manuel_prioritaire_sur_lestimation_automatique(self, dossier_config_isole):
        from logic.rapports import participants_recus_effectifs
        from logic.configuration import charger_configuration, sauvegarder_configuration

        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls"),  # estimation auto = 1
        ])
        config = charger_configuration()
        config["participants_recus_manuel"] = {"Tchad": 8}  # correction manuelle de l'agent
        sauvegarder_configuration(config)

        recus = participants_recus_effectifs(evt_df)
        assert recus["Tchad"] == 8  # la valeur manuelle l'emporte, pas 1

    def test_sans_correction_manuelle_lestimation_automatique_est_utilisee(self, dossier_config_isole):
        from logic.rapports import participants_recus_effectifs
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/B/f.xls"),
        ])
        recus = participants_recus_effectifs(evt_df)
        assert recus["Tchad"] == 2

    def test_generation_du_rapport_reflete_les_attendus_personnalises(self, dossier_config_isole):
        from logic.configuration import charger_configuration, sauvegarder_configuration
        from docx import Document

        config = charger_configuration()
        config["participants_attendus"] = {"Tchad": 100}
        sauvegarder_configuration(config)

        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        assert erreur is None
        doc = Document(chemin)
        ligne_tchad = [r for r in doc.tables[0].rows if r.cells[0].text == "Tchad"][0]
        assert ligne_tchad.cells[1].text == "100"
        os.remove(chemin)


class TestGraphiques:
    """
    Comparaison (taux de reponse par pays) et evolution (score par mois,
    calculee sur les donnees ACTUELLEMENT traitees -- pas l'historique)
    -- incorpores au rapport d'activite sous forme d'images, avec une
    explication textuelle adaptee au contenu reel, pas un texte fixe.
    """

    def test_comparaison_genere_un_fichier_image_valide(self, dossier_config_isole):
        from logic.rapports import _graphique_comparaison_pays
        import os as os_module
        taux = [
            {"pays": "Tchad", "attendu": 12, "recu": 8, "pct": 66.7},
            {"pays": "Cameroun", "attendu": 23, "recu": 2, "pct": 8.7},
            {"pays": "TOTAL", "attendu": 35, "recu": 10, "pct": 28.6},
        ]
        chemin = _graphique_comparaison_pays(taux)
        assert chemin is not None
        assert os_module.path.exists(chemin)
        assert os_module.path.getsize(chemin) > 0
        os_module.remove(chemin)

    def test_comparaison_ignore_la_ligne_total(self, dossier_config_isole):
        """La ligne TOTAL ne doit jamais etre tracee comme un pays."""
        from logic.rapports import _graphique_comparaison_pays
        taux_uniquement_total = [{"pays": "TOTAL", "attendu": 10, "recu": 5, "pct": 50.0}]
        assert _graphique_comparaison_pays(taux_uniquement_total) is None

    def test_evolution_absente_si_un_seul_mois(self, dossier_config_isole):
        """Un seul mois present dans les donnees -> rien a comparer."""
        from logic.rapports import _graphique_evolution_pays
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril")])
        assert _graphique_evolution_pays(evt_df, None) is None

    def test_evolution_generee_avec_plusieurs_mois(self, dossier_config_isole, mapping_minimal_pret):
        import os as os_module
        from logic.rapports import _graphique_evolution_pays
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril"),
            _evt("Tchad", "Tchad/MOIS FEVRIER 2026/B/f.xls", mois="Février"),
        ])
        chemin = _graphique_evolution_pays(evt_df, mapping_minimal_pret)
        assert chemin is not None
        assert os_module.path.exists(chemin)
        os_module.remove(chemin)

    def test_evolution_avec_plusieurs_pays_et_mois(self, dossier_config_isole, mapping_minimal_pret):
        """Cas reel decrit : plusieurs pays valides ensemble, chacun avec
        plusieurs mois -- une courbe par pays, sans planter."""
        import os as os_module
        from logic.rapports import _graphique_evolution_pays
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril"),
            _evt("Tchad", "Tchad/MOIS FEVRIER 2026/B/f.xls", mois="Février"),
            _evt("Cameroun", "Cameroun/MOIS AVRIL 2026/C/f.xls", mois="Avril"),
        ])
        chemin = _graphique_evolution_pays(evt_df, mapping_minimal_pret)
        assert chemin is not None
        os_module.remove(chemin)

    def test_rapport_activite_contient_deux_images_avec_plusieurs_mois(self, dossier_config_isole, mapping_minimal_pret):
        from docx import Document

        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril"),
            _evt("Tchad", "Tchad/MOIS FEVRIER 2026/B/f.xls", mois="Février"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, mapping_minimal_pret)
        assert erreur is None

        doc = Document(chemin)
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        nb_images = sum(
            1 for p in doc.paragraphs for run in p.runs
            if run._element.findall(f".//{ns}blip")
        )
        # 4 images attendues : logo (page de garde) + comparaison + evolution + statut
        assert nb_images == 4
        os.remove(chemin)


class TestGenerationRapportActivite:

    def test_genere_un_fichier_valide(self, tmp_path, monkeypatch, dossier_config_isole):
        monkeypatch.chdir(tmp_path.parent if False else ".")  # no-op, garde le cwd du projet
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls", statut="R"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        assert erreur is None
        assert chemin is not None
        assert os.path.exists(chemin)
        os.remove(chemin)

    def test_tableau_taux_de_reponse_contient_les_six_pays_plus_total(self, tmp_path, dossier_config_isole):
        from docx import Document
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        assert erreur is None
        doc = Document(chemin)
        table = doc.tables[0]
        assert len(table.rows) == 8  # entete + 6 pays + TOTAL
        pays_dans_tableau = [row.cells[0].text for row in table.rows[1:-1]]
        assert set(pays_dans_tableau) == set(participants_attendus_par_pays().keys())
        assert table.rows[-1].cells[0].text == "TOTAL"
        os.remove(chemin)

    def test_taux_faible_declenche_les_mesures_incitatrices(self, dossier_config_isole):
        from docx import Document
        # Un seul participant recu sur 60 attendus -> taux tres faible
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        doc = Document(chemin)
        textes = " ".join(p.text for p in doc.paragraphs)
        assert "mesures incitatrices" in textes.lower()
        assert "insignifiant" in textes.lower()
        os.remove(chemin)

    def test_taux_eleve_necrit_pas_les_mesures_incitatrices(self, dossier_config_isole):
        from docx import Document
        lignes = []
        for pays, attendu in participants_attendus_par_pays().items():
            for i in range(attendu):
                lignes.append(_evt(pays, f"{pays}/MOIS AVRIL 2026/Banque_{i}/f.xls", survenu=False))
        evt_df = pd.DataFrame(lignes)
        chemin, erreur = generer_rapport_activite(evt_df, None)
        doc = Document(chemin)
        textes = " ".join(p.text for p in doc.paragraphs)
        assert "mesures incitatrices" not in textes.lower()
        assert "satisfaisant" in textes.lower()
        os.remove(chemin)

    def test_statut_correctement_categorise_dans_le_tableau_analyse(self, dossier_config_isole):
        from docx import Document
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", statut="R"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/B/f.xls", statut="ER"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/C/f.xls", statut="NR"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        doc = Document(chemin)
        table_analyse = doc.tables[1]
        ligne_tchad = [r for r in table_analyse.rows if r.cells[0].text == "Tchad"][0]
        # colonnes : PAYS | Nb | % | Resolu | En cours | Non resolu
        assert ligne_tchad.cells[3].text == "1"  # Resolu (R)
        assert ligne_tchad.cells[4].text == "1"  # En cours (ER)
        assert ligne_tchad.cells[5].text == "1"  # Non resolu (NR)
        os.remove(chemin)

    def test_remarques_citent_le_taux_reellement_calcule(self, dossier_config_isole):
        from docx import Document
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        doc = Document(chemin)
        textes = " ".join(p.text for p in doc.paragraphs)
        # 1 participant recu sur 60 attendus = 1.7%
        assert "1.7 %" in textes
        os.remove(chemin)

    def test_template_manquant_renvoie_une_erreur_explicite(self, monkeypatch, dossier_config_isole):
        import logic.rapports as rap
        chemin_original = "templates/rapport_activite.docx"
        vrai_exists = os.path.exists
        monkeypatch.setattr(rap.os.path, "exists", lambda p: False if str(p).replace("\\", "/").endswith(chemin_original) else vrai_exists(p))
        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/BANQUE_A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        assert chemin is None
        assert erreur is not None


class TestModificationTableauParticipants:
    """
    L'agent doit pouvoir ajouter ou retirer des pays dans le tableau
    Attendu/Reçu de Paramètres (st.data_editor num_rows="dynamic") --
    ces tests verifient la fonction d'extraction independamment du
    widget Streamlit, qui n'est pas simulable par les outils de test.
    """

    def test_ajout_dun_nouveau_pays(self):
        from logic.configuration import construire_participants_depuis_tableau
        df = pd.DataFrame([
            {"Pays": "Cameroun", "Attendu": 23, "Reçu (estimation automatique)": 5,
             "Reçu (correction manuelle, optionnel)": None},
            {"Pays": "Nigeria", "Attendu": 40, "Reçu (estimation automatique)": 0,
             "Reçu (correction manuelle, optionnel)": None},
        ])
        attendus, _ = construire_participants_depuis_tableau(df)
        assert attendus["Nigeria"] == 40

    def test_retrait_dun_pays(self):
        from logic.configuration import construire_participants_depuis_tableau
        df = pd.DataFrame([
            {"Pays": "Cameroun", "Attendu": 23, "Reçu (estimation automatique)": 5,
             "Reçu (correction manuelle, optionnel)": None},
        ])
        attendus, _ = construire_participants_depuis_tableau(df)
        assert set(attendus.keys()) == {"Cameroun"}
        assert "Tchad" not in attendus

    def test_ligne_sans_nom_de_pays_est_ignoree(self):
        from logic.configuration import construire_participants_depuis_tableau
        df = pd.DataFrame([
            {"Pays": "Cameroun", "Attendu": 23, "Reçu (estimation automatique)": 5,
             "Reçu (correction manuelle, optionnel)": None},
            {"Pays": "", "Attendu": None, "Reçu (estimation automatique)": None,
             "Reçu (correction manuelle, optionnel)": None},
        ])
        attendus, _ = construire_participants_depuis_tableau(df)
        assert len(attendus) == 1

    def test_tableau_vide_produit_une_configuration_vide(self):
        from logic.configuration import construire_participants_depuis_tableau
        df = pd.DataFrame(columns=["Pays", "Attendu", "Reçu (estimation automatique)", "Reçu (correction manuelle, optionnel)"])
        attendus, recus = construire_participants_depuis_tableau(df)
        assert attendus == {}
        assert recus == {}

    def test_pays_ajoute_apparait_ensuite_dans_le_rapport(self, dossier_config_isole):
        """Un pays ajoute par l'agent (hors des 6 CEMAC habituels) doit
        se retrouver dans le tableau du rapport d'activite genere apres
        enregistrement."""
        from logic.configuration import construire_participants_depuis_tableau, charger_configuration, sauvegarder_configuration
        from docx import Document

        df = pd.DataFrame([
            {"Pays": "Nigeria", "Attendu": 15, "Reçu (estimation automatique)": 0,
             "Reçu (correction manuelle, optionnel)": None},
        ])
        attendus, recus = construire_participants_depuis_tableau(df)
        config = charger_configuration()
        config["participants_attendus"] = attendus
        config["participants_recus_manuel"] = recus
        sauvegarder_configuration(config)

        evt_df = pd.DataFrame([_evt("Nigeria", "Nigeria/MOIS AVRIL 2026/Banque/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, None)
        assert erreur is None
        doc = Document(chemin)
        pays_dans_tableau = [row.cells[0].text for row in doc.tables[0].rows[1:-1]]
        assert "Nigeria" in pays_dans_tableau
        assert "Tchad" not in pays_dans_tableau  # retire, ne doit plus apparaitre
        os.remove(chemin)


class TestEspacementEtImages:
    """
    Le nettoyage de l'espacement (paragraphes vides excedentaires laisses
    par les balises Jinja apres rendu) ne doit JAMAIS supprimer un
    paragraphe qui contient une image (graphique, logo) -- son .text est
    vide comme un vrai paragraphe vide, mais ce n'en est pas un.
    """

    def test_nettoyage_preserve_les_images(self):
        from logic.rapports import nettoyer_espacement
        from docx import Document
        from docx.shared import Mm

        chemin_logo = "referentiels/logo.png"
        if not os.path.exists(chemin_logo):
            pytest.skip("Logo indisponible dans cet environnement")

        doc = Document()
        doc.add_paragraph("Titre")
        doc.add_paragraph("")
        doc.add_paragraph("")
        p_image = doc.add_paragraph()
        run = p_image.add_run()
        run.add_picture(chemin_logo, width=Mm(10))
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("Texte final")

        nb_images_avant = len(doc.inline_shapes)
        nettoyer_espacement(doc)
        nb_images_apres = len(doc.inline_shapes)

        assert nb_images_avant == nb_images_apres == 1

    def test_nettoyage_reduit_bien_les_sequences_de_texte_vide(self):
        from logic.rapports import nettoyer_espacement
        from docx import Document

        doc = Document()
        doc.add_paragraph("Titre")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("Texte final")

        nettoyer_espacement(doc)
        vides = [p for p in doc.paragraphs if p.text.strip() == ""]
        assert len(vides) == 1

    def test_rapport_active_genere_conserve_toutes_ses_images(self, dossier_config_isole, mapping_minimal_pret):
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril"),
            _evt("Tchad", "Tchad/MOIS FEVRIER 2026/B/f.xls", mois="Février"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, mapping_minimal_pret)
        assert erreur is None
        from docx import Document
        doc = Document(chemin)
        # logo (page de garde) + comparaison + evolution (2 mois presents) + statut = 4
        assert len(doc.inline_shapes) == 4
        os.remove(chemin)


class TestApercuCorrespondAExport:
    """
    L'apercu HTML affiche dans l'application doit toujours contenir les
    memes images (logo, graphiques) que le document Word reellement
    exporte -- un paragraphe porteur d'une image a un texte vide, ce qui
    l'avait fait ignorer a tort par le generateur d'apercu.
    """

    def test_apercu_contient_autant_dimages_que_lexport(self, dossier_config_isole, mapping_minimal_pret):
        from logic.rapports import generer_apercu_html
        from docx import Document

        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", mois="Avril"),
            _evt("Tchad", "Tchad/MOIS FEVRIER 2026/B/f.xls", mois="Février"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, mapping_minimal_pret)
        assert erreur is None

        doc = Document(chemin)
        nb_images_export = len(doc.inline_shapes)

        apercu_html = generer_apercu_html(chemin)
        nb_images_apercu = apercu_html.count("<img")

        assert nb_images_apercu == nb_images_export
        assert nb_images_export > 0  # verifie que le test porte bien sur un cas avec images
        os.remove(chemin)

    def test_apercu_ne_contient_pas_de_paragraphe_vide_pour_une_image(self, dossier_config_isole, mapping_minimal_pret):
        """Une image ne doit jamais etre rendue comme un paragraphe texte
        vide (invisible a l'ecran) -- elle doit produire une vraie balise img."""
        from logic.rapports import generer_apercu_html

        evt_df = pd.DataFrame([_evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls")])
        chemin, erreur = generer_rapport_activite(evt_df, mapping_minimal_pret)
        assert erreur is None

        apercu_html = generer_apercu_html(chemin)
        assert "data:image/png;base64," in apercu_html
        os.remove(chemin)


class TestGraphiqueStatutDysfonctionnements:
    """
    Graphique ajoute pour la section '1.2 Analyse des reportings' :
    repartition Resolu/En cours/Non resolu, calculee en sommant le
    tableau analyse_pays_ctx sur tous les pays.
    """

    def test_genere_un_fichier_quand_des_dysfonctionnements_existent(self, dossier_config_isole):
        from logic.rapports import _graphique_statut_dysfonctionnements
        analyse = [
            {"pays": "Tchad", "nb": 17, "pct": 94.4, "resolu": 5, "en_cours": 6, "non_resolu": 6},
            {"pays": "Cameroun", "nb": 1, "pct": 5.6, "resolu": 1, "en_cours": 0, "non_resolu": 0},
        ]
        chemin = _graphique_statut_dysfonctionnements(analyse)
        assert chemin is not None
        assert os.path.exists(chemin)
        os.remove(chemin)

    def test_absent_si_aucun_dysfonctionnement(self, dossier_config_isole):
        from logic.rapports import _graphique_statut_dysfonctionnements
        analyse = [{"pays": "Tchad", "nb": 0, "pct": 0, "resolu": 0, "en_cours": 0, "non_resolu": 0}]
        assert _graphique_statut_dysfonctionnements(analyse) is None

    def test_fonctionne_avec_une_seule_categorie_non_nulle(self, dossier_config_isole):
        """Si tout est Resolu (ou toute autre categorie unique), le
        camembert doit quand meme se generer sans planter (une seule
        tranche a 100%)."""
        from logic.rapports import _graphique_statut_dysfonctionnements
        analyse = [{"pays": "Tchad", "nb": 5, "pct": 100.0, "resolu": 5, "en_cours": 0, "non_resolu": 0}]
        chemin = _graphique_statut_dysfonctionnements(analyse)
        assert chemin is not None
        os.remove(chemin)

    def test_rapport_activite_integre_le_graphique_de_statut(self, dossier_config_isole, mapping_minimal_pret):
        from docx import Document
        evt_df = pd.DataFrame([
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/A/f.xls", statut="R"),
            _evt("Tchad", "Tchad/MOIS AVRIL 2026/B/f.xls", statut="NR"),
        ])
        chemin, erreur = generer_rapport_activite(evt_df, mapping_minimal_pret)
        assert erreur is None
        doc = Document(chemin)
        textes = " ".join(p.text for p in doc.paragraphs)
        assert "répartition des" in textes and "dysfonctionnements" in textes
        os.remove(chemin)
