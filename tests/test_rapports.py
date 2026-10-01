# -*- coding: utf-8 -*-
"""
Tests de logic.rapports.generer_pdf_natif : generation PDF en Python pur
(reportlab), garantie de fonctionner sans Microsoft Word ni LibreOffice
installes -- notamment sur Windows, ou l'alternative weasyprint aurait
introduit une dependance systeme externe (GTK3) tout aussi genante que
celle qu'on cherchait a eviter.
"""
import os
from docx import Document

from logic.rapports import generer_pdf_natif


def _construire_docx_complexe(chemin):
    """Document couvrant les cas de rendu principaux : titres, texte,
    puces, tableau -- dans cet ordre, pour verifier que l'ordre reel du
    document est respecte (pas seulement son contenu)."""
    doc = Document()
    doc.add_heading("Titre principal", level=1)
    doc.add_paragraph("Un paragraphe normal.")
    doc.add_paragraph("Premier point", style="List Bullet")
    doc.add_paragraph("Second point", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Colonne A"
    table.rows[0].cells[1].text = "Colonne B"
    table.rows[1].cells[0].text = "val1"
    table.rows[1].cells[1].text = "val2"
    doc.add_paragraph("Texte après le tableau.")
    doc.save(chemin)


class TestGenerationPdfNative:

    def test_produit_un_fichier_pdf_valide(self, tmp_path):
        chemin_docx = str(tmp_path / "test.docx")
        _construire_docx_complexe(chemin_docx)

        chemin_pdf = generer_pdf_natif(chemin_docx)

        assert chemin_pdf is not None
        assert os.path.exists(chemin_pdf)
        assert chemin_pdf.endswith(".pdf")
        with open(chemin_pdf, "rb") as f:
            entete = f.read(5)
        assert entete == b"%PDF-"  # signature de fichier PDF standard

    def test_document_vide_ne_plante_pas(self, tmp_path):
        chemin_docx = str(tmp_path / "vide.docx")
        Document().save(chemin_docx)
        resultat = generer_pdf_natif(chemin_docx)
        assert resultat is None  # rien a rendre, mais pas d'exception

    def test_caracteres_speciaux_ne_cassent_pas_le_rendu(self, tmp_path):
        """Les caracteres XML/HTML speciaux (&, <, >) doivent etre
        echappes plutot que casser la generation."""
        chemin_docx = str(tmp_path / "special.docx")
        doc = Document()
        doc.add_paragraph("Test avec & un < signe > et \"guillemets\"")
        doc.save(chemin_docx)

        chemin_pdf = generer_pdf_natif(chemin_docx)
        assert chemin_pdf is not None
        assert os.path.exists(chemin_pdf)

    def test_gere_un_document_avec_uniquement_un_tableau(self, tmp_path):
        chemin_docx = str(tmp_path / "tableau_seul.docx")
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        table.rows[0].cells[2].text = "C"
        doc.save(chemin_docx)

        chemin_pdf = generer_pdf_natif(chemin_docx)
        assert chemin_pdf is not None
        assert os.path.exists(chemin_pdf)
