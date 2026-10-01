# -*- coding: utf-8 -*-
"""
Tests de logic.rapports.generer_rapport_general : recapitulatif compact
(tableau, une ligne par incident) de tous les incidents survenus, tous
pays confondus, trie par pays puis par mois.
"""
import os
import pandas as pd
import pytest

from logic.rapports import generer_rapport_general


def _evt(pays, mois="Avril", domaine="SYGMA", reference="RPSYGPS01",
         statut="R", occurrence=None, duree=None, motif=None):
    return {
        "reference": reference, "evenement": "Test", "survenance": "Oui",
        "occurrence": occurrence, "motif": motif, "statut": statut, "duree_indispo_min": duree,
        "commentaire": None, "domaine": domaine, "site": "PS",
        "fichier_source": f"{pays}/MOIS {mois.upper()} 2026/Banque/{reference}.xlsx",
        "detection_confiance": 1.0, "detection_avertissement": None, "hors_referentiel": False,
        "survenance_bin": 1, "occurrence_num": occurrence,
        "pays": pays, "annee": "2026", "mois": mois,
    }


class TestFormatCompact:

    def test_genere_un_seul_tableau(self, tmp_path):
        evt_df = pd.DataFrame([_evt("Tchad"), _evt("Cameroun", reference="RPSWIPS01", domaine="SWIFT")])
        chemin, erreur = generer_rapport_general(evt_df, None)
        assert erreur is None
        from docx import Document
        doc = Document(chemin)
        assert len(doc.tables) == 1
        os.remove(chemin)

    def test_une_ligne_par_incident_tous_pays_confondus(self):
        """Le tableau doit contenir directement tous les pays, sans
        section separee par pays a parcourir."""
        from docx import Document
        evt_df = pd.DataFrame([
            _evt("Tchad", reference="RPSYGPS01"),
            _evt("Cameroun", reference="RPSWIPS01", domaine="SWIFT"),
            _evt("Congo", reference="RPSYSPS01", domaine="SYSTAC"),
        ])
        chemin, erreur = generer_rapport_general(evt_df, None)
        doc = Document(chemin)
        table = doc.tables[0]
        pays_presents = {row.cells[0].text for row in table.rows[1:]}
        assert pays_presents == {"Tchad", "Cameroun", "Congo"}
        os.remove(chemin)

    def test_trie_par_pays_puis_par_mois(self):
        from docx import Document
        evt_df = pd.DataFrame([
            _evt("Tchad", mois="Mai", reference="RPSYGPS02"),
            _evt("Tchad", mois="Avril", reference="RPSYGPS01"),
            _evt("Cameroun", mois="Avril", reference="RPSWIPS01", domaine="SWIFT"),
        ])
        chemin, erreur = generer_rapport_general(evt_df, None)
        doc = Document(chemin)
        table = doc.tables[0]
        lignes = [(row.cells[0].text, row.cells[1].text) for row in table.rows[1:]]
        # Cameroun avant Tchad (ordre alphabetique des pays), et au sein de
        # Tchad, Avril avant Mai.
        assert lignes == [("Cameroun", "Avril"), ("Tchad", "Avril"), ("Tchad", "Mai")]
        os.remove(chemin)

    def test_champs_manquants_affichent_un_repli_explicite_pas_une_erreur(self):
        from docx import Document
        evt_df = pd.DataFrame([_evt("Tchad", occurrence=None, duree=None, motif=None)])
        chemin, erreur = generer_rapport_general(evt_df, None)
        assert erreur is None
        doc = Document(chemin)
        ligne = doc.tables[0].rows[1]
        valeurs = [c.text for c in ligne.cells]
        assert "(non renseigné)" in valeurs or "(non renseignée)" in valeurs
        os.remove(chemin)

    def test_document_est_en_orientation_paysage(self):
        """Format paysage retenu pour maximiser le nombre de colonnes par
        ligne et minimiser le nombre de pages."""
        from docx import Document
        from docx.enum.section import WD_ORIENT
        evt_df = pd.DataFrame([_evt("Tchad")])
        chemin, erreur = generer_rapport_general(evt_df, None)
        doc = Document(chemin)
        assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE
        os.remove(chemin)

    def test_aucun_incident_survenu_produit_un_tableau_vide_sans_planter(self):
        from docx import Document
        evt_df = pd.DataFrame([{**_evt("Tchad"), "survenance_bin": 0}])
        chemin, erreur = generer_rapport_general(evt_df, None)
        assert erreur is None
        doc = Document(chemin)
        assert len(doc.tables[0].rows) == 1  # seulement l'entete

    def test_template_manquant_renvoie_une_erreur_explicite(self, monkeypatch):
        import logic.rapports as rap
        chemin_original = "templates/rapport_general.docx"
        vrai_exists = os.path.exists
        monkeypatch.setattr(rap.os.path, "exists", lambda p: False if str(p).replace("\\", "/").endswith(chemin_original) else vrai_exists(p))
        evt_df = pd.DataFrame([_evt("Tchad")])
        chemin, erreur = generer_rapport_general(evt_df, None)
        assert chemin is None
        assert erreur is not None
