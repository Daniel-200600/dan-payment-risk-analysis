# -*- coding: utf-8 -*-
"""
Fixtures partagees pour les tests DAN.

Principe general : chaque test qui touche a un fichier sur disque (historique,
mapping, guide de surveillance...) redirige les constantes de chemin des
modules concernes vers un dossier temporaire (via monkeypatch), afin de ne
jamais lire ni ecrire dans les vrais fichiers du projet pendant les tests.
"""
import io
import shutil
import pandas as pd
import pytest
from docx import Document


class FichierUploadeFactice(io.BytesIO):
    """
    Reproduit l'interface minimale d'un fichier retourne par
    st.file_uploader (attributs .name et .size) a partir d'octets bruts,
    pour pouvoir tester lire_reporting() sans dependre de Streamlit.
    """
    def __init__(self, donnees, nom):
        super().__init__(donnees)
        self.name = nom
        self.size = len(donnees)


@pytest.fixture
def fabriquer_excel():
    """
    Fabrique un faux fichier de reporting Excel a partir d'une liste de
    lignes (dont la premiere doit etre l'en-tete). Renvoie un objet
    directement passable a lire_reporting().
    """
    import openpyxl

    def _fabriquer(lignes, nom="test.xlsx"):
        wb = openpyxl.Workbook()
        ws = wb.active
        for ligne in lignes:
            ws.append(ligne)
        buf = io.BytesIO()
        wb.save(buf)
        return FichierUploadeFactice(buf.getvalue(), nom)

    return _fabriquer


@pytest.fixture
def entete_standard():
    """En-tete standard d'un fichier de reporting DAN, dans l'ordre attendu."""
    return ["REFERENCE", "EVENEMENT", "SURVENANCE", "OCCURRENCE",
            "MOTIF", "STATUT", "DUREE", "COMMENTAIRE"]


@pytest.fixture
def mapping_minimal(tmp_path):
    """
    Petite table de correspondance autonome (5 lignes), suffisante pour
    tester le calcul de score sans dependre du mapping complet
    du projet. Renvoie le DataFrame directement (pas un chemin de fichier).
    """
    return pd.DataFrame({
        "Code événement": ["RPSYGPS01", "RPSWIPS01", "RPSYSPS01", "RPSYSPS02", "RPRESPS01"],
        "Événement": ["Evt SYGMA", "Evt SWIFT", "Evt SYSTAC 1", "Evt SYSTAC 2", "Evt RESEAU"],
        "Domaine": ["SYGMA", "SWIFT", "SYSTAC", "SYSTAC", "RESEAU"],
        "Site": ["PS", "PS", "PS", "PA", "PS"],
        "Code référentiel": ["RS1", "RS2", "RS3", "RS4", "RS5"],
        "Élément d'appréciation": ["a", "b", "c", "d", "e"],
        "Risque(s)": ["RO", "RJ", "RO+RJ", "RL", "RO"],
        "Niveau": ["Critique", "Majeur", "Sensible", "Faible", "Critique"],
        "Confiance": ["Forte", "Forte", "Moyenne", "Forte", "Faible"],
        "Justification": ["j1", "j2", "j3", "j4", "j5"],
    })


@pytest.fixture
def mapping_minimal_pret(mapping_minimal):
    """
    Meme contenu que mapping_minimal, mais deja renomme comme le fait
    charger_mapping() (colonnes en minuscules : reference, risque,
    niveau...). A utiliser pour tester calculer_scores() directement,
    qui attend ce format -- mapping_minimal seul reste au format brut
    (colonnes Excel d'origine) pour tester la validation/le remplacement
    du fichier source.
    """
    return mapping_minimal.rename(columns={
        "Code événement": "reference", "Événement": "evenement_ref",
        "Domaine": "domaine", "Site": "site",
        "Code référentiel": "code_referentiel",
        "Élément d'appréciation": "element_appreciation",
        "Risque(s)": "risque", "Niveau": "niveau", "Confiance": "confiance",
        "Justification": "justification",
    })


@pytest.fixture
def dossier_historique_isole(tmp_path, monkeypatch):
    """Redirige logic.historique vers un fichier temporaire isole."""
    import logic.historique as hist
    chemin = tmp_path / "historique.json"
    monkeypatch.setattr(hist, "CHEMIN_HISTORIQUE", str(chemin))
    return chemin


@pytest.fixture
def dossier_referentiel_isole(tmp_path, monkeypatch, mapping_minimal):
    """
    Redirige logic.referentiel (table de correspondance ET guide de
    surveillance) vers des fichiers/dossiers temporaires isoles, avec un
    mapping et un guide de depart deja en place.
    """
    import logic.referentiel as ref

    chemin_mapping = tmp_path / "mapping_risques.xlsx"
    mapping_minimal.to_excel(chemin_mapping, index=False)
    monkeypatch.setattr(ref, "CHEMIN_MAPPING", str(chemin_mapping))
    monkeypatch.setattr(ref, "DOSSIER_VERSIONS", str(tmp_path / "versions_mapping"))

    chemin_guide = tmp_path / "GUIDE_DE_SURVEILLANCE.docx"
    doc = Document()
    doc.add_paragraph("Guide de test - version initiale")
    doc.save(chemin_guide)
    monkeypatch.setattr(ref, "CHEMIN_GUIDE", str(chemin_guide))
    monkeypatch.setattr(ref, "DOSSIER_VERSIONS_GUIDE", str(tmp_path / "versions_guide"))

    return tmp_path


@pytest.fixture
def dossier_config_isole(tmp_path, monkeypatch):
    """Redirige logic.configuration vers un fichier temporaire isole."""
    import logic.configuration as conf
    chemin = tmp_path / "configuration.json"
    monkeypatch.setattr(conf, "CHEMIN_CONFIG", str(chemin))
    return chemin
