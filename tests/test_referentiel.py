# -*- coding: utf-8 -*-
"""
Tests de logic.referentiel : validation d'un fichier candidat, remplacement
avec sauvegarde, et surtout non-collision entre deux operations rapprochees
-- un bug reel a corrompu le vrai Guide de surveillance pendant le
developpement (deux sauvegardes horodatees a la seconde pres se sont
ecrasees l'une l'autre lors d'un remplacement suivi d'une restauration
executes dans la meme seconde).
"""
import io
import pandas as pd
import pytest
from docx import Document


class TestValidationMapping:

    def test_rejette_fichier_avec_colonnes_manquantes(self, dossier_referentiel_isole):
        import logic.referentiel as ref
        df_incomplet = pd.DataFrame({"Code événement": ["RPSYGPS99"], "Événement": ["Test"]})
        buf = io.BytesIO()
        df_incomplet.to_excel(buf, index=False)
        buf.seek(0)
        valide, messages, _ = ref.valider_fichier_mapping(buf)
        assert not valide

    def test_accepte_fichier_complet(self, dossier_referentiel_isole, mapping_minimal):
        import logic.referentiel as ref
        buf = io.BytesIO()
        mapping_minimal.to_excel(buf, index=False)
        buf.seek(0)
        valide, messages, df = ref.valider_fichier_mapping(buf)
        assert valide
        assert len(df) == len(mapping_minimal)

    def test_domaine_vide_est_reconstruit_depuis_le_code(self, dossier_referentiel_isole):
        """Protection contre la recurrence du bug historique ou la colonne
        Domaine restait entierement vide dans le fichier source."""
        import logic.referentiel as ref
        df = pd.DataFrame({
            "Code événement": ["RPSYSPS01"], "Événement": ["Test"], "Domaine": [None],
            "Site": ["PS"], "Code référentiel": ["RS1"], "Élément d'appréciation": ["x"],
            "Risque(s)": ["RO"], "Niveau": ["Faible"], "Confiance": ["Forte"], "Justification": ["j"],
        })
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)
        valide, messages, df_corrige = ref.valider_fichier_mapping(buf)
        assert valide
        assert df_corrige["Domaine"].iloc[0] == "SYSTAC"

    def test_domaine_incoherent_est_corrige_depuis_le_code(self, dossier_referentiel_isole):
        import logic.referentiel as ref
        df = pd.DataFrame({
            "Code événement": ["RPSYGPS01"], "Événement": ["Test"], "Domaine": ["SWIFT"],  # faux
            "Site": ["PS"], "Code référentiel": ["RS1"], "Élément d'appréciation": ["x"],
            "Risque(s)": ["RO"], "Niveau": ["Faible"], "Confiance": ["Forte"], "Justification": ["j"],
        })
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)
        _, _, df_corrige = ref.valider_fichier_mapping(buf)
        assert df_corrige["Domaine"].iloc[0] == "SYGMA"  # le code (RPSYG) fait foi


class TestRemplacementEtRestaurationMapping:

    def test_remplacement_sauvegarde_lancienne_version(self, dossier_referentiel_isole, mapping_minimal):
        import logic.referentiel as ref
        nouveau = mapping_minimal.copy()
        nouveau = nouveau.iloc[:2]  # version plus courte
        ref.remplacer_mapping(nouveau)
        assert len(ref.resume_mapping()["par_domaine"]) <= 2
        assert len(ref.lister_versions_precedentes()) == 1

    def test_restauration_redonne_le_contenu_dorigine(self, dossier_referentiel_isole, mapping_minimal):
        import logic.referentiel as ref
        etat_avant = ref.resume_mapping()["nb_lignes"]

        nouveau = mapping_minimal.iloc[:1].copy()
        ref.remplacer_mapping(nouveau)
        assert ref.resume_mapping()["nb_lignes"] == 1

        versions = ref.lister_versions_precedentes()
        ref.restaurer_version(versions[0])
        assert ref.resume_mapping()["nb_lignes"] == etat_avant

    def test_operations_rapprochees_ne_se_corrompent_pas(self, dossier_referentiel_isole, mapping_minimal):
        """Reproduit le bug reel : remplacer puis restaurer immediatement
        (dans la meme fraction de seconde) ne doit plus faire perdre le
        contenu original -- l'horodatage des sauvegardes doit etre assez
        precis pour ne jamais entrer en collision."""
        import logic.referentiel as ref
        etat_avant = ref.resume_mapping()["nb_lignes"]

        nouveau = mapping_minimal.iloc[:1].copy()
        ref.remplacer_mapping(nouveau)
        versions = ref.lister_versions_precedentes()
        ref.restaurer_version(versions[0])  # execute juste apres, sans delai

        assert ref.resume_mapping()["nb_lignes"] == etat_avant


class TestGuideDeSurveillance:

    def test_rejette_fichier_non_docx(self, dossier_referentiel_isole):
        import logic.referentiel as ref

        class Faux(io.BytesIO):
            name = "guide.pdf"

        valide, message = ref.valider_fichier_guide(Faux(b"contenu"))
        assert not valide

    def test_accepte_document_word_valide(self, dossier_referentiel_isole):
        import logic.referentiel as ref
        doc = Document()
        doc.add_paragraph("Contenu de test")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "guide.docx"
        valide, message = ref.valider_fichier_guide(buf)
        assert valide

    def test_remplacement_et_restauration_du_guide(self, dossier_referentiel_isole):
        import logic.referentiel as ref
        nb_paragraphes_avant = ref.resume_guide()["nb_paragraphes"]

        doc = Document()
        doc.add_paragraph("Nouveau contenu")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "nouveau_guide.docx"
        ref.remplacer_guide(buf)
        assert ref.resume_guide()["nb_paragraphes"] == 1

        versions = ref.lister_versions_precedentes_guide()
        ref.restaurer_version_guide(versions[0])
        assert ref.resume_guide()["nb_paragraphes"] == nb_paragraphes_avant

    def test_operations_guide_rapprochees_ne_se_corrompent_pas(self, dossier_referentiel_isole):
        """Meme regression que pour le mapping, mais sur le Guide de
        surveillance -- c'est ce cas precis qui a corrompu le vrai
        document pendant le developpement avant correction."""
        import logic.referentiel as ref
        nb_paragraphes_avant = ref.resume_guide()["nb_paragraphes"]

        doc = Document()
        doc.add_paragraph("A")
        doc.add_paragraph("B")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "test.docx"
        ref.remplacer_guide(buf)

        versions = ref.lister_versions_precedentes_guide()
        ref.restaurer_version_guide(versions[0])

        assert ref.resume_guide()["nb_paragraphes"] == nb_paragraphes_avant
