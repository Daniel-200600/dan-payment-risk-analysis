# -*- coding: utf-8 -*-
"""
Generate the SYNTHETIC resources shipped with the public DAN repository.

Everything produced here is invented for demonstration purposes:

  mapping/mapping_risques.xlsx        illustrative event -> risk mapping
  referentiels/GUIDE_DE_SURVEILLANCE.docx   short generic user guide
  referentiels/logo.png               placeholder logo
  templates/*.docx                    docxtpl report templates (generic wording)
  data/sample/                        fictitious monthly reporting forms

Run from the repository root:

    python scripts/generate_demo_resources.py

The output is deterministic (fixed random seed).
"""
import os
import random

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

BLEU = RGBColor(0x0B, 0x2E, 0x59)

# --------------------------------------------------------------------------
# 1. Synthetic risk mapping
# --------------------------------------------------------------------------
# (domain code, domain label, site, [(event label, risk, level, confidence)])
EVENEMENTS = {
    ("SYG", "SYGMA", "PS"): [
        ("Interruption de la liaison avec la plate-forme centrale", "RO", "Critique", "Forte"),
        ("Indisponibilité d'un composant technique du participant", "RO", "Majeur", "Forte"),
        ("Expiration d'un certificat électronique", "RO+RJ", "Sensible", "Moyenne"),
        ("Rejet de messages par la plate-forme", "RO", "Sensible", "Forte"),
        ("Réclamation d'un client sur un virement", "RJ", "A valider", "Faible"),
        ("Basculement vers le site de secours", "RO", "Majeur", "Moyenne"),
    ],
    ("SWI", "SWIFT", "PS"): [
        ("Panne d'un périphérique d'impression", "RO", "Faible", "Forte"),
        ("Perte de connexion entre un poste et le serveur local", "RO", "Sensible", "Forte"),
        ("Déconnexion de la messagerie interbancaire", "RO", "Critique", "Forte"),
        ("Ajout d'un compte utilisateur", "RJ", "Sensible", "Moyenne"),
        ("Message refusé par le réseau", "RO", "Sensible", "Forte"),
        ("Échec d'une opération de sauvegarde", "RO+RL", "A valider", "Faible"),
    ],
    ("SYS", "SYSTAC", "PS"): [
        ("Indisponibilité du poste de capture au siège", "RO", "Sensible", "Forte"),
        ("Échec de transmission d'un lot d'opérations", "RO", "Majeur", "Moyenne"),
        ("Anomalie de rapprochement des opérations", "RO+RJ", "A valider", "Moyenne"),
        ("Retard de traitement en fin de journée", "RO", "Sensible", "Forte"),
        ("Dysfonctionnement de l'application de capture", "RO", "Majeur", "Forte"),
        ("Incident de sauvegarde quotidienne", "RO", "Sensible", "Moyenne"),
    ],
    ("SYS", "SYSTAC", "PA"): [
        ("Indisponibilité du poste de capture en agence", "RO", "Sensible", "Forte"),
        ("Rupture de liaison entre l'agence et le siège", "RO", "Majeur", "Forte"),
        ("Rejet d'un lot d'opérations en agence", "RO", "A valider", "Faible"),
        ("Erreur de saisie détectée avant transmission", "RJ", "Faible", "Moyenne"),
    ],
    ("RES", "RESEAU", "PS"): [
        ("Coupure de la liaison réseau principale", "RO", "Critique", "Forte"),
        ("Basculement sur la liaison de secours", "RO", "Majeur", "Moyenne"),
        ("Dégradation de la bande passante", "RO", "Sensible", "Moyenne"),
        ("Alerte de sécurité sur un équipement réseau", "RO+RJ", "A valider", "Faible"),
        ("Maintenance réseau non planifiée", "RO", "Sensible", "Forte"),
    ],
}


def construire_mapping():
    lignes = []
    for (code, domaine, site), evts in EVENEMENTS.items():
        for i, (libelle, risque, niveau, confiance) in enumerate(evts, start=1):
            lignes.append({
                "Code événement": f"RP{code}{site}{i:02d}",
                "Événement": libelle,
                "Domaine": domaine,
                "Site": site,
                "Code référentiel": f"REF-{code}-{site}-{i:02d}",
                "Élément d'appréciation": "Élément illustratif (données de démonstration)",
                "Risque(s)": risque,
                "Niveau": niveau,
                "Confiance": confiance,
                "Justification": "Correspondance fictive, fournie à titre d'exemple.",
            })
    return pd.DataFrame(lignes)


# --------------------------------------------------------------------------
# 2. Placeholder logo
# --------------------------------------------------------------------------
def creer_logo(chemin):
    img = Image.new("RGBA", (360, 360), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((10, 10, 350, 350), radius=48, fill=(11, 46, 89, 255))
    d.rounded_rectangle((28, 28, 332, 332), radius=36, outline=(201, 162, 39, 255), width=8)
    try:
        police = ImageFont.truetype("arialbd.ttf", 120)
    except OSError:
        police = ImageFont.load_default()
    d.text((180, 180), "DAN", font=police, fill=(255, 255, 255, 255), anchor="mm")
    img.save(chemin)


# --------------------------------------------------------------------------
# 3. Generic user guide
# --------------------------------------------------------------------------
def creer_guide(chemin):
    doc = Document()
    doc.add_heading("Guide de surveillance (version de démonstration)", level=1)
    doc.add_paragraph(
        "Ce document est un exemple générique fourni avec le dépôt public de DAN. "
        "Il ne reprend aucun document réglementaire ou interne : remplacez-le par le "
        "guide de votre propre organisation (Paramètres > Guide de surveillance)."
    )
    doc.add_heading("Principe", level=2)
    doc.add_paragraph(
        "Chaque participant transmet un reporting mensuel listant des événements "
        "(survenus ou non). DAN rapproche ces événements d'une table de correspondance "
        "de risques, calcule un score de gravité et produit des rapports."
    )
    doc.add_heading("Codification des références", level=2)
    doc.add_paragraph(
        "Format : RP + domaine (SYG, SWI, SYS, RES) + site (PS ou PA) + numéro à deux "
        "chiffres. Exemple : RPSYGPS01."
    )
    doc.save(chemin)


# --------------------------------------------------------------------------
# 4. Report templates (docxtpl)
# --------------------------------------------------------------------------
def _para(doc, texte, gras=False, taille=None, centre=False, couleur=None):
    p = doc.add_paragraph()
    r = p.add_run(texte)
    r.bold = gras
    if taille:
        r.font.size = Pt(taille)
    if couleur is not None:
        r.font.color.rgb = couleur
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _pied(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.text = ("Document généré automatiquement par DAN — outil d'aide à la surveillance "
              "des systèmes de paiement (données de démonstration).")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(8)


def _tableau_boucle(doc, entetes, ligne_for, cellules, ligne_endfor="{%tr endfor %}"):
    """Table with docxtpl row-loop markers on their own rows."""
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    for i, e in enumerate(entetes):
        t.rows[0].cells[i].text = e
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    ligne = t.add_row().cells
    ligne[0].text = ligne_for
    donnees = t.add_row().cells
    for i, c in enumerate(cellules):
        donnees[i].text = c
    fin = t.add_row().cells
    fin[0].text = ligne_endfor
    return t


DESCRIPTIONS = {
    "SYGMA": "Système SYGMA — règlement de gros montants",
    "SWIFT": "Système SWIFT — messagerie interbancaire",
    "SYSTAC": "Système SYSTAC — télécompensation",
    "RESEAU": "Système RESEAU — infrastructure réseau",
}


def creer_modele_systeme(chemin, domaine):
    doc = Document()
    _para(doc, "RAPPORT DE SURVEILLANCE MENSUEL", gras=True, taille=18, centre=True, couleur=BLEU)
    _para(doc, DESCRIPTIONS[domaine], taille=12, centre=True)
    _para(doc, "{{ logo }}", centre=True)
    _para(doc, "Participant : {{ participant }}")
    _para(doc, "Fichier source : {{ fichier_source }}")
    _para(doc, "Période : {{ mois }} / {{ annee }}")
    _para(doc, "1. Événements survenus ce mois-ci", gras=True, taille=13, couleur=BLEU)
    _para(doc, "{%p if nb_evenements_survenus and nb_evenements_survenus > 0 %}")
    _para(doc, "{%p for e in evenements if e.survenu %}")
    _para(doc, "{{ e.reference }} — {{ e.evenement }}", gras=True)
    _para(doc, "Motif : {{ e.motif }}")
    _para(doc, "Statut : {{ e.statut }}")
    _para(doc, "Occurrences ce mois-ci : {{ e.occurrence }}")
    _para(doc, "Durée d'indisponibilité : {{ e.duree }}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p else %}")
    _para(doc, "Aucun événement n'est survenu sur ce système ce mois-ci.")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if synthese %}")
    _para(doc, "2. Synthèse qualitative", gras=True, taille=13, couleur=BLEU)
    _para(doc, "{{ synthese }}")
    _para(doc, "{%p endif %}")
    _para(doc, "3. Recommandations", gras=True, taille=13, couleur=BLEU)
    _para(doc, "{{ recommandations }}")
    _para(doc, "Rapport {} généré automatiquement par DAN le {{{{ date_generation }}}}.".format(domaine), taille=9)
    _pied(doc)
    doc.save(chemin)


def creer_modele_general(chemin, logo):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    _para(doc, "RAPPORT GÉNÉRAL DES INCIDENTS", gras=True, taille=18, centre=True, couleur=BLEU)
    _para(doc, "Participant : {{ participant }}")
    _para(doc, "Période : {{ mois }} {{ annee }}")
    _para(doc, "Récapitulatif de tous les incidents survenus, tous pays confondus", gras=True, taille=13, couleur=BLEU)
    _para(doc, "{{ nb_incidents_total }} incident(s) survenu(s), sur {{ nb_pays }} pays.")
    _tableau_boucle(
        doc,
        ["Pays", "Mois", "Système", "Référence", "Événement", "Motif", "Statut", "Occ.", "Durée"],
        "{%tr for i in incidents_ctx %}",
        ["{{ i.pays }}", "{{ i.mois }}", "{{ i.domaine }}", "{{ i.reference }}", "{{ i.evenement }}",
         "{{ i.motif }}", "{{ i.statut }}", "{{ i.occurrence }}", "{{ i.duree }}"],
    )
    _para(doc, "Généré le {{ date_generation }}.", taille=9)
    _pied(doc)
    doc.save(chemin)


def creer_modele_activite(chemin, logo):
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(logo, width=Mm(30))
    _para(doc, "RAPPORT D'ACTIVITÉ", gras=True, taille=22, centre=True, couleur=BLEU)
    _para(doc, "Surveillance des systèmes de paiement", taille=13, centre=True)
    _para(doc, "Période : {{ mois }} {{ annee }}", centre=True)
    _para(doc, "Généré le {{ date_generation }}", centre=True, taille=9)

    _para(doc, "Sommaire", gras=True, taille=14, couleur=BLEU)
    _para(doc, "Chapitre I : Présentation générale")
    _para(doc, "Chapitre II : Activités de la surveillance")

    _para(doc, "Chapitre I : Présentation générale", gras=True, taille=14, couleur=BLEU)
    _para(doc, "Cette section est un texte générique à adapter à votre organisation. La "
               "surveillance des systèmes de paiement a pour objet de s'assurer de leur bon "
               "fonctionnement, de leur sécurité et de leur efficacité, à partir des "
               "reportings transmis périodiquement par les participants.")
    _para(doc, "Les systèmes suivis sont : SYGMA (règlement de gros montants), SYSTAC "
               "(télécompensation), SWIFT (messagerie interbancaire) et l'infrastructure RESEAU.")

    _para(doc, "Chapitre II : Activités de la surveillance", gras=True, taille=14, couleur=BLEU)
    _para(doc, "1. Analyse des reportings", gras=True, taille=12, couleur=BLEU)
    _para(doc, "1.1 Taux de réponse par pays", gras=True)
    _tableau_boucle(doc, ["Pays", "Attendu", "Reçu", "Taux (%)"], "{%tr for t in taux_reponse_ctx %}",
                    ["{{ t.pays }}", "{{ t.attendu }}", "{{ t.recu }}", "{{ t.pct }}"])
    _para(doc, "{{ synthese_taux_reponse }}")
    _para(doc, "{%p if mesures_incitatrices %}")
    _para(doc, "Afin d'améliorer ce taux, les mesures incitatrices suivantes sont envisagées :")
    _para(doc, "{%p for m in mesures_incitatrices %}")
    _para(doc, "• {{ m }}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if graphique_comparaison %}")
    _para(doc, "{{ graphique_comparaison }}", centre=True)
    _para(doc, "{{ explication_comparaison }}")
    _para(doc, "{%p endif %}")

    _para(doc, "1.2 Analyse des reportings par pays", gras=True)
    _tableau_boucle(doc, ["Pays", "Nb", "%", "Résolu", "En cours", "Non résolu"],
                    "{%tr for a in analyse_pays_ctx %}",
                    ["{{ a.pays }}", "{{ a.nb }}", "{{ a.pct }}", "{{ a.resolu }}",
                     "{{ a.en_cours }}", "{{ a.non_resolu }}"])
    _para(doc, "{{ synthese_dysfonctionnements }}")
    _para(doc, "{%p if graphique_statut %}")
    _para(doc, "{{ graphique_statut }}", centre=True)
    _para(doc, "{{ explication_statut }}")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if graphique_evolution %}")
    _para(doc, "{{ graphique_evolution }}", centre=True)
    _para(doc, "{{ explication_evolution }}")
    _para(doc, "{%p endif %}")

    _para(doc, "2. Contrôle sur place", gras=True, taille=12, couleur=BLEU)
    _para(doc, "(Section à compléter par l'analyste.)")
    _para(doc, "3. Autres activités", gras=True, taille=12, couleur=BLEU)
    _para(doc, "(Section à compléter par l'analyste.)")
    _para(doc, "4. Difficultés rencontrées", gras=True, taille=12, couleur=BLEU)
    _para(doc, "{%p if difficultes %}")
    _para(doc, "{%p for d in difficultes %}")
    _para(doc, "• {{ d }}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p else %}")
    _para(doc, "Aucune difficulté particulière à signaler.")
    _para(doc, "{%p endif %}")
    _para(doc, "5. Remarques", gras=True, taille=12, couleur=BLEU)
    _para(doc, "{{ remarques_activite }}")
    _pied(doc)
    doc.save(chemin)


# --------------------------------------------------------------------------
# 5. Fictitious monthly reporting forms
# --------------------------------------------------------------------------
MOTIFS = [
    "Incident technique ponctuel", "Maintenance non planifiée", "Coupure d'alimentation électrique",
    "Lenteur de la liaison opérateur", "Erreur de configuration corrigée",
]
ENTETE = ["REFERENCE", "EVENEMENT", "SURVENANCE (Oui/Non)", "NOMBRE D'OCCURRENCE",
          "MOTIF/CAUSE", "STATUT DES EVENEMENTS (R=Résolu; ER=Encours de résolution; NR=Non Résolu)",
          "DUREE D'INDISPONIBILITE (en minute)", "COMMENTAIRE DE L'ASSUJETTI"]


def ecrire_formulaire(chemin, titre, evts, rng):
    wb = Workbook()
    ws = wb.active
    ws["A3"] = f"REPORTING DE SURVEILLANCE {titre} (données fictives)"
    # Year/month are left blank on purpose: DAN reads them from the folder
    # names (e.g. "MOIS MARS 2026") and falls back on these cells otherwise.
    ws["D6"] = "Année"
    ws["D7"] = "Mois"
    ws["D8"] = "Code Banque"
    for j, h in enumerate(ENTETE, start=1):
        ws.cell(row=11, column=j, value=h)
    for i, (code, libelle) in enumerate(evts, start=12):
        survenu = rng.random() < 0.25
        ws.cell(row=i, column=1, value=code)
        ws.cell(row=i, column=2, value=libelle)
        ws.cell(row=i, column=3, value="Oui" if survenu else "Non")
        if survenu:
            ws.cell(row=i, column=4, value=rng.randint(1, 6))
            ws.cell(row=i, column=5, value=rng.choice(MOTIFS))
            ws.cell(row=i, column=6, value=rng.choice(["R", "ER", "NR"]))
            ws.cell(row=i, column=7, value=rng.choice([0, 15, 30, 45, 90]))
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    wb.save(chemin)


def creer_donnees_exemple(racine, mapping):
    rng = random.Random(2026)
    for pays in ("Cameroun", "Gabon", "Tchad"):
        for banque in ("BANQUE_A", "BANQUE_B"):
            for domaine, groupe in mapping.groupby("Domaine"):
                evts = list(zip(groupe["Code événement"], groupe["Événement"]))
                dossier = os.path.join(racine, pays, "MOIS MARS 2026", banque)
                ecrire_formulaire(os.path.join(dossier, f"{domaine}_PARTICIPANT.xlsx"),
                                  f"{domaine} PARTICIPANT", evts, rng)


# --------------------------------------------------------------------------
def main():
    for d in ("mapping", "referentiels", "templates", "data/sample"):
        os.makedirs(d, exist_ok=True)
    mapping = construire_mapping()
    mapping.to_excel("mapping/mapping_risques.xlsx", index=False)
    creer_logo("referentiels/logo.png")
    creer_guide("referentiels/GUIDE_DE_SURVEILLANCE.docx")
    for dom in DESCRIPTIONS:
        creer_modele_systeme(f"templates/rapport_{dom}.docx", dom)
    creer_modele_general("templates/rapport_general.docx", "referentiels/logo.png")
    creer_modele_activite("templates/rapport_activite.docx", "referentiels/logo.png")
    creer_donnees_exemple("data/sample", mapping)
    print("Synthetic demo resources generated.")


if __name__ == "__main__":
    main()
